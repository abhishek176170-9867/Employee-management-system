import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    NAVY = RGBColor(30, 27, 75)        # #1E1B4B (Primary Dark)
    INDIGO = RGBColor(79, 70, 229)     # #4F46E5 (Primary Accent)
    VIOLET = RGBColor(124, 58, 237)    # #7C3AED (Secondary Accent)
    EMERALD = RGBColor(5, 150, 105)    # #059669 (Success)
    AMBER = RGBColor(217, 119, 6)      # #D97706 (Warning)
    DARK_TEXT = RGBColor(15, 23, 42)   # #0F172A
    MUTED_TEXT = RGBColor(100, 116, 139) # #64748B
    CARD_BG = RGBColor(248, 250, 252)  # #F8FAFC
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0
    WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title, category="LEAVEFLOW ENTERPRISE"):
        # Header Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.8), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(238, 242, 255)
        badge.line.color.rgb = RGBColor(199, 210, 254)
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.text = f"●  {category.upper()}"
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = INDIGO
        p_b.alignment = PP_ALIGN.CENTER

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_TEXT

        # Accent Rule Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(226, 232, 240)
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg=CARD_BG, border=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg
        card.line.color.rgb = border
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Hero Dark Navy)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Decorative Top Accent Bar
    accent_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = INDIGO
    accent_bar.line.fill.background()

    # Project Pill Badge
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.3), Inches(3.2), Inches(0.42))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(49, 46, 129)
    pill.line.color.rgb = INDIGO
    tf_p = pill.text_frame
    p_p = tf_p.paragraphs[0]
    p_p.text = "ENTERPRISE SOFTWARE PROJECT"
    p_p.font.size = Pt(11)
    p_p.font.bold = True
    p_p.font.color.rgb = RGBColor(165, 180, 252)
    p_p.alignment = PP_ALIGN.CENTER

    # Main Project Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Employee Leave Management System"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "LeaveFlow — Next-Generation Automated Corporate Time-Off & Attendance Platform"
    p1_sub.font.size = Pt(18)
    p1_sub.font.color.rgb = RGBColor(199, 210, 254)
    p1_sub.space_before = Pt(12)

    # Key Highlights Grid (3 Mini Badges)
    highlights = [
        ("Smart Working Days", "Automated exclusion of weekends & public holidays"),
        ("Dual-Role Portals", "Employee Self-Service & HR Administrator Queue"),
        ("100% Tested Architecture", "17/17 Automated Unit Tests with Zero-Defect QA")
    ]
    for i, (title, desc) in enumerate(highlights):
        c = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0 + i * 3.8), Inches(4.2), Inches(3.6), Inches(1.2))
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(49, 46, 129)
        c.line.color.rgb = RGBColor(79, 70, 229)
        c_tf = c.text_frame
        c_tf.word_wrap = True
        c_tf.margin_left = c_tf.margin_top = Inches(0.2)
        cp1 = c_tf.paragraphs[0]
        cp1.text = f"✔ {title}"
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = WHITE
        cp2 = c_tf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(10.5)
        cp2.font.color.rgb = RGBColor(199, 210, 254)
        cp2.space_before = Pt(4)

    # Footer Metadata Card
    f_card = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.3), Inches(13.333), Inches(1.2))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = RGBColor(17, 24, 39)
    f_card.line.fill.background()

    f_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.45), Inches(11.3), Inches(0.9))
    f_tf = f_box.text_frame
    f_p = f_tf.paragraphs[0]
    f_p.text = "Student Presenter: Abhishek Yadav    |    Roll No: 06    |    Academic Year: 2026"
    f_p.font.size = Pt(14)
    f_p.font.bold = True
    f_p.font.color.rgb = WHITE
    f_p2 = f_tf.add_paragraph()
    f_p2.text = "Full-Stack Technology Stack: Python 3.12 • Flask 3.0 • SQLite3 • Tailwind CSS • Chart.js • Render Cloud"
    f_p2.font.size = Pt(11)
    f_p2.font.color.rgb = RGBColor(156, 163, 175)
    f_p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 2: Project Overview & Problem Statement
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Project Overview & Business Problem", "Executive Summary")

    # Left Column: Problem Statement Card
    add_card(s2, Inches(0.8), Inches(1.9), Inches(5.6), Inches(5.0), bg=RGBColor(254, 242, 242), border=RGBColor(254, 202, 202))
    pb = s2.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.6))
    ptf = pb.text_frame
    ptf.word_wrap = True
    p = ptf.paragraphs[0]
    p.text = "The Legacy Challenge: Manual & Paper Leaves"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(185, 28, 28)

    problems = [
        ("Calculation Errors:", "Weekends and company public holidays manually included, unfairly deducting employee allowances."),
        ("Approval Delays:", "Physical paper forms and unorganized emails cause multi-day review bottlenecks for managers."),
        ("Quota Discrepancies:", "Lack of automated pending holds leads to double-booking and negative quota balances."),
        ("Audit Non-Compliance:", "No centralized immutable audit logs or CSV compliance exports for HR audits.")
    ]
    for title, desc in problems:
        p_item = ptf.add_paragraph()
        p_item.text = f"❌  {title} {desc}"
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = DARK_TEXT
        p_item.space_before = Pt(12)

    # Right Column: The LeaveFlow Solution Card
    add_card(s2, Inches(6.8), Inches(1.9), Inches(5.7), Inches(5.0), bg=RGBColor(240, 253, 244), border=RGBColor(187, 247, 208))
    sb = s2.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.6))
    stf = sb.text_frame
    stf.word_wrap = True
    p_s = stf.paragraphs[0]
    p_s.text = "The LeaveFlow Solution: Automated Cloud System"
    p_s.font.size = Pt(16)
    p_s.font.bold = True
    p_s.font.color.rgb = EMERALD

    solutions = [
        ("Smart Calendar Engine:", "Dynamic arithmetic excluding Saturdays, Sundays, and gazetted company holidays automatically."),
        ("Instant Self-Service:", "Real-time balance computation, half-day leaves, and immediate cancellation with quota refund."),
        ("HR Approvals Queue:", "One-click manager review modal with custom feedback notes and live attendance tracking."),
        ("Compliance & Analytics:", "Interactive Chart.js visualizations, automated audit trails, and multi-filter CSV streaming.")
    ]
    for title, desc in solutions:
        s_item = stf.add_paragraph()
        s_item.text = f"✔  {title} {desc}"
        s_item.font.size = Pt(12)
        s_item.font.color.rgb = DARK_TEXT
        s_item.space_before = Pt(12)

    # =========================================================================
    # SLIDE 3: Project Objectives
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Project Strategic Objectives", "Core Goals & Scope")

    objectives = [
        ("1. End-to-End Lifecycle Automation",
         "Transform employee time-off requests into a paperless digital workflow from submission through approval, quota deduction, and historical archiving.",
         INDIGO),
        ("2. Intelligent Working-Days Engine",
         "Eliminate human error by automatically filtering company-observed holidays and weekends from leave calculations in real-time.",
         EMERALD),
        ("3. Strict Role-Based Access Control",
         "Isolate employee self-service actions from administrative privileges, ensuring role-based privacy and secure session management.",
         VIOLET),
        ("4. Real-Time Analytics & Audit Compliance",
         "Empower HR leaders with dynamic Chart.js utilization trends, employee distribution metrics, and audit-ready CSV exports.",
         AMBER)
    ]

    for i, (title, desc, color) in enumerate(objectives):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.9 + row * 2.5)

        card = add_card(s3, left, top, Inches(5.7), Inches(2.2))
        
        # Color bar indicator on left
        bar = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(0.18), Inches(2.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = s3.shapes.add_textbox(left + Inches(0.4), top + Inches(0.25), Inches(5.0), Inches(1.7))
        tframe = tb.text_frame
        tframe.word_wrap = True
        tp = tframe.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(15)
        tp.font.bold = True
        tp.font.color.rgb = DARK_TEXT

        tp2 = tframe.add_paragraph()
        tp2.text = desc
        tp2.font.size = Pt(12)
        tp2.font.color.rgb = MUTED_TEXT
        tp2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 4: Technology Stack
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Enterprise Technology Stack", "Architecture & Tools")

    tech_pillars = [
        ("Backend Services", INDIGO, [
            ("Python 3.12+", "High-performance programming language powering business computation"),
            ("Flask 3.0", "Micro-web framework providing WSGI routing, templating, and REST APIs"),
            ("Werkzeug Security", "PBKDF2 SHA-256 cryptographic password hashing"),
            ("Gunicorn WSGI", "Production-grade UNIX WSGI HTTP server for scalable cloud serving")
        ]),
        ("Frontend & UI/UX", VIOLET, [
            ("Tailwind CSS", "Modern utility-first CSS framework for clean, responsive UI layouts"),
            ("Vanilla JS (ES6+)", "Client-side interactive calculator, live duration badge, modal controls"),
            ("Chart.js", "Hardware-accelerated HTML5 canvas charts for management analytics"),
            ("FontAwesome 6", "Vector iconography for intuitive visual cues and status tags")
        ]),
        ("Database & DevOps", EMERALD, [
            ("SQLite 3 Engine", "Zero-configuration ACID transactional database with PRAGMA FK checks"),
            ("Python unittest", "Comprehensive 17-test automated verification test harness"),
            ("Git & GitHub", "Version control, tracking, and continuous deployment synchronization"),
            ("Render Cloud Platform", "Containerized cloud deployment with automated SSL/HTTPS hosting")
        ])
    ]

    for i, (pillar, color, items) in enumerate(tech_pillars):
        left = Inches(0.8 + i * 3.95)
        card = add_card(s4, left, Inches(1.9), Inches(3.8), Inches(5.0))
        
        # Header strip on card
        strip = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.9), Inches(3.8), Inches(0.7))
        strip.fill.solid()
        strip.fill.fore_color.rgb = color
        strip.line.fill.background()
        st_tf = strip.text_frame
        st_p = st_tf.paragraphs[0]
        st_p.text = pillar
        st_p.font.size = Pt(14)
        st_p.font.bold = True
        st_p.font.color.rgb = WHITE
        st_p.alignment = PP_ALIGN.CENTER

        tb = s4.shapes.add_textbox(left + Inches(0.2), Inches(2.75), Inches(3.4), Inches(4.0))
        tframe = tb.text_frame
        tframe.word_wrap = True

        for k, (name, d) in enumerate(items):
            p = tframe.add_paragraph() if k > 0 else tframe.paragraphs[0]
            p.text = f"•  {name}"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = DARK_TEXT
            p.space_before = Pt(8) if k > 0 else Pt(0)

            p_sub = tframe.add_paragraph()
            p_sub.text = d
            p_sub.font.size = Pt(10.5)
            p_sub.font.color.rgb = MUTED_TEXT
            p_sub.space_before = Pt(2)

    # =========================================================================
    # SLIDE 5: System Architecture & Database Design
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "System Architecture & Database Schema", "Structural Design")

    # Left: Architecture Flow Box
    add_card(s5, Inches(0.8), Inches(1.9), Inches(6.0), Inches(5.0))
    tb_arch = s5.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.6), Inches(4.6))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True
    p = tf_arch.paragraphs[0]
    p.text = "Multi-Tier Software Architecture"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INDIGO

    tiers = [
        ("Presentation Tier (Client):", "Responsive Web Browser, Tailwind CSS, FontAwesome, Chart.js, asynchronous JSON REST calls (`/api/calculate-days`)."),
        ("Application Tier (Flask Backend):", "Authentication controller, RBAC gatekeepers, Business Calendar Day Engine, and CSV streaming pipeline."),
        ("Persistence Tier (SQLite3 DB):", "Relational database enforcing referential foreign key constraints and ACID transactional updates.")
    ]
    for title, desc in tiers:
        pt = tf_arch.add_paragraph()
        pt.text = f"🏛️  {title}"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = DARK_TEXT
        pt.space_before = Pt(10)

        pd = tf_arch.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(2)

    # Right: Database Entities
    add_card(s5, Inches(7.1), Inches(1.9), Inches(5.4), Inches(5.0))
    tb_db = s5.shapes.add_textbox(Inches(7.3), Inches(2.1), Inches(5.0), Inches(4.6))
    tf_db = tb_db.text_frame
    tf_db.word_wrap = True
    p_db = tf_db.paragraphs[0]
    p_db.text = "Relational Database Schema (5 Tables)"
    p_db.font.size = Pt(16)
    p_db.font.bold = True
    p_db.font.color.rgb = EMERALD

    tables = [
        ("users", "Employee records, hashed passwords, department, role (admin / employee)"),
        ("leave_types", "Annual, Sick, Casual, Maternity, Paternity, Unpaid categories"),
        ("leave_balances", "Yearly allocated, used, and pending days per user & category"),
        ("leave_applications", "Start/end dates, total_days, half_day flag, remarks, review state"),
        ("audit_logs", "Immutable audit trail of all logins, submissions, and approvals")
    ]
    for tbl, desc in tables:
        pt = tf_db.add_paragraph()
        pt.text = f"📋  {tbl}"
        pt.font.size = Pt(12.5)
        pt.font.bold = True
        pt.font.color.rgb = DARK_TEXT
        pt.space_before = Pt(8)

        pd = tf_db.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(1)

    # =========================================================================
    # SLIDE 6: Key Features — Employee Self-Service
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Employee Self-Service Experience", "Feature Walkthrough")

    emp_features = [
        ("Smart Duration Calculator",
         "Live calculation badge instantly computes working business days as the user selects dates. Saturdays, Sundays, and official public holidays are automatically subtracted.",
         INDIGO),
        ("Single-Day Half-Day Support",
         "Toggling 'Half-Day' automatically synchronizes end date to start date, recording exactly 0.5 days without manual overrides.",
         EMERALD),
        ("Real-Time Quota Snapshot",
         "Color-coded progress gauges display total allocated, used, under review, and remaining balance for each leave category.",
         VIOLET),
        ("Instant Self-Service Cancellation",
         "Employees can cancel pending leave requests with one click, automatically releasing the pending hold back to available balance.",
         AMBER)
    ]

    for i, (title, desc, color) in enumerate(emp_features):
        left = Inches(0.8 + (i % 2) * 5.9)
        top = Inches(1.9 + (i // 2) * 2.5)
        add_card(s6, left, top, Inches(5.7), Inches(2.2))

        b = s6.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.3), top + Inches(0.35), Inches(0.4), Inches(0.4))
        b.fill.solid()
        b.fill.fore_color.rgb = color
        b.line.fill.background()

        tb = s6.shapes.add_textbox(left + Inches(0.9), top + Inches(0.25), Inches(4.5), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tp = tf.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(15)
        tp.font.bold = True
        tp.font.color.rgb = DARK_TEXT

        tp2 = tf.add_paragraph()
        tp2.text = desc
        tp2.font.size = Pt(11.5)
        tp2.font.color.rgb = MUTED_TEXT
        tp2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 7: Key Features — HR & Management Administration
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "HR Administration & Analytics Portal", "Management Control")

    admin_features = [
        ("Approvals Queue Workflow",
         "Centralized queue displaying pending applications with employee details, date spans, and remaining quotas. Fast-track modal sign-off with custom manager feedback remarks.",
         INDIGO),
        ("Staff & Quota Calibration",
         "HR administrators can register new employees with automatic baseline quota initialization, and modify annual allowances dynamically on demand.",
         EMERALD),
        ("Executive Chart.js Analytics",
         "Interactive visual dashboard featuring donut distribution charts by leave category and monthly utilization bar trends across departments.",
         VIOLET),
        ("Compliance Search & CSV Export",
         "Multi-dimensional search filters (by department, status, and category) with one-click export to audit-ready CSV spreadsheets for legal/HR records.",
         AMBER)
    ]

    for i, (title, desc, color) in enumerate(admin_features):
        left = Inches(0.8 + (i % 2) * 5.9)
        top = Inches(1.9 + (i // 2) * 2.5)
        add_card(s7, left, top, Inches(5.7), Inches(2.2))

        b = s7.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.3), top + Inches(0.35), Inches(0.4), Inches(0.4))
        b.fill.solid()
        b.fill.fore_color.rgb = color
        b.line.fill.background()

        tb = s7.shapes.add_textbox(left + Inches(0.9), top + Inches(0.25), Inches(4.5), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tp = tf.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(15)
        tp.font.bold = True
        tp.font.color.rgb = DARK_TEXT

        tp2 = tf.add_paragraph()
        tp2.text = desc
        tp2.font.size = Pt(11.5)
        tp2.font.color.rgb = MUTED_TEXT
        tp2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 8: System Verification & Quality Assurance
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Verification & Quality Assurance (100% Passed)", "Testing & Reliability")

    # Big Metric Banner
    metric_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.733), Inches(1.2))
    metric_card.fill.solid()
    metric_card.fill.fore_color.rgb = RGBColor(240, 253, 244)
    metric_card.line.color.rgb = RGBColor(187, 247, 208)
    metric_card.line.width = Pt(1)

    m_tb = s8.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(11.1), Inches(0.9))
    m_tf = m_tb.text_frame
    m_p = m_tf.paragraphs[0]
    m_p.text = "✔  Automated Test Suite: 17 of 17 Tests Passed (100% Pass Rate in 4.44s)"
    m_p.font.size = Pt(18)
    m_p.font.bold = True
    m_p.font.color.rgb = EMERALD

    m_p2 = m_tf.add_paragraph()
    m_p2.text = "Comprehensive automated unit and integration tests executing against isolated temporary SQLite databases on every build."
    m_p2.font.size = Pt(12)
    m_p2.font.color.rgb = DARK_TEXT
    m_p2.space_before = Pt(4)

    # 3 Sub-Cards for Test Categories
    test_categories = [
        ("Calculation & Engine Tests", [
            ("Regular Business Days:", "Mon-Wed correctly computed as 3 days"),
            ("Weekend Exclusion:", "Fri-Mon computes 2 days (Sat/Sun skipped)"),
            ("Holiday Exclusion:", "Public holidays excluded from count"),
            ("Half-Day Validation:", "Exact 0.5-day duration verified"),
            ("Negative Date Orders:", "End date < start date rejected")
        ]),
        ("Workflow & State Tests", [
            ("Pending Hold Creation:", "Applying locks quota to prevent double-booking"),
            ("Manager Approval:", "Transfers pending hold to used days"),
            ("Manager Rejection:", "Releases pending hold with no penalty"),
            ("Employee Cancellation:", "Instant restoration of held quota"),
            ("Insufficient Balance:", "Rejects request when quota exceeded")
        ]),
        ("Security & Edge-Case Tests", [
            ("Role-Based Protection:", "Prevents unauthorized admin access"),
            ("Overlapping Collisions:", "Rejects conflicting concurrent leaves"),
            ("Email Uniqueness:", "Prevents duplicate employee registration"),
            ("Negative Quota Guards:", "Validates allocation thresholds"),
            ("Audit Trail Verification:", "Verifies log entry on all critical operations")
        ])
    ]

    for i, (cat_title, test_items) in enumerate(test_categories):
        left = Inches(0.8 + i * 3.95)
        add_card(s8, left, Inches(3.3), Inches(3.8), Inches(3.6))

        tb = s8.shapes.add_textbox(left + Inches(0.2), Inches(3.45), Inches(3.4), Inches(3.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tp = tf.paragraphs[0]
        tp.text = cat_title
        tp.font.size = Pt(13)
        tp.font.bold = True
        tp.font.color.rgb = INDIGO

        for item_name, item_desc in test_items:
            p = tf.add_paragraph()
            p.text = f"• {item_name} {item_desc}"
            p.font.size = Pt(10)
            p.font.color.rgb = DARK_TEXT
            p.space_before = Pt(5)

    # =========================================================================
    # SLIDE 9: Conclusion & Project Achievements
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Conclusion & Project Achievements", "Summary & Roadmap")

    # Left: Achievements Card
    add_card(s9, Inches(0.8), Inches(1.9), Inches(5.7), Inches(5.0))
    tb_c = s9.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.1), Inches(4.6))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Key Project Achievements"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INDIGO

    achievements = [
        ("Enterprise-Ready Architecture:", "Delivered a fully responsive, zero-defect employee leave platform with modern aesthetics."),
        ("100% Automated Calculations:", "Eliminated manual computation errors via dynamic holiday-aware date algorithms."),
        ("Robust Security & Compliance:", "PBKDF2 SHA-256 password security, role isolation, and comprehensive audit logs."),
        ("Cloud Production Deployment:", "Configured with Gunicorn WSGI and Render cloud blueprints for instant zero-downtime serving.")
    ]
    for title, desc in achievements:
        pt = tf_c.add_paragraph()
        pt.text = f"⭐  {title}"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = DARK_TEXT
        pt.space_before = Pt(12)

        pd = tf_c.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(2)

    # Right: Future Scope & Roadmap
    add_card(s9, Inches(6.8), Inches(1.9), Inches(5.7), Inches(5.0))
    tb_f = s9.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.6))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = "Future Enhancement Scope"
    p_f.font.size = Pt(16)
    p_f.font.bold = True
    p_f.font.color.rgb = EMERALD

    roadmap = [
        ("Instant Slack & Teams Webhooks:", "Real-time notifications sent to company messaging channels on application submission and approvals."),
        ("Two-Way Google & Outlook Calendar Sync:", "Approved leaves automatically reflected on team shared Google Calendars."),
        ("Biometric & Punch Clock Attendance:", "Integration with hardware fingerprint/facial attendance scanners."),
        ("AI Predictive Absence Forecasting:", "Machine learning models forecasting seasonal departmental staffing deficits.")
    ]
    for title, desc in roadmap:
        pt = tf_f.add_paragraph()
        pt.text = f"🚀  {title}"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = DARK_TEXT
        pt.space_before = Pt(12)

        pd = tf_f.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(2)

    # =========================================================================
    # SLIDE 10: Thank You & Q&A
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = NAVY
    bg10.line.fill.background()

    # Top accent line
    top_line = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = EMERALD
    top_line.line.fill.background()

    tb10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    p10 = tf10.paragraphs[0]
    p10.text = "Thank You!"
    p10.font.size = Pt(48)
    p10.font.bold = True
    p10.font.color.rgb = WHITE
    p10.alignment = PP_ALIGN.CENTER

    p10_sub = tf10.add_paragraph()
    p10_sub.text = "Employee Leave Management System (LeaveFlow)"
    p10_sub.font.size = Pt(22)
    p10_sub.font.bold = True
    p10_sub.font.color.rgb = RGBColor(165, 180, 252)
    p10_sub.alignment = PP_ALIGN.CENTER
    p10_sub.space_before = Pt(10)

    p10_qa = tf10.add_paragraph()
    p10_qa.text = "Questions & Answers"
    p10_qa.font.size = Pt(18)
    p10_qa.font.color.rgb = RGBColor(209, 213, 219)
    p10_qa.alignment = PP_ALIGN.CENTER
    p10_qa.space_before = Pt(16)

    # Contact / Repository Box
    card10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(4.5), Inches(8.333), Inches(1.8))
    card10.fill.solid()
    card10.fill.fore_color.rgb = RGBColor(49, 46, 129)
    card10.line.color.rgb = INDIGO
    c10_tf = card10.text_frame
    c10_tf.word_wrap = True

    p_repo = c10_tf.paragraphs[0]
    p_repo.text = "Project Metadata & Repository"
    p_repo.font.size = Pt(13)
    p_repo.font.bold = True
    p_repo.font.color.rgb = RGBColor(199, 210, 254)
    p_repo.alignment = PP_ALIGN.CENTER

    p_info = c10_tf.add_paragraph()
    p_info.text = "Presenter: Abhishek Yadav  |  Roll No: 06  |  Email: abhishek176170@gmail.com"
    p_info.font.size = Pt(13)
    p_info.font.color.rgb = WHITE
    p_info.alignment = PP_ALIGN.CENTER
    p_info.space_before = Pt(6)

    p_link = c10_tf.add_paragraph()
    p_link.text = "GitHub Repository: https://github.com/abhishek176170-9867/Employee-management-system"
    p_link.font.size = Pt(12)
    p_link.font.bold = True
    p_link.font.color.rgb = RGBColor(52, 211, 153)
    p_link.alignment = PP_ALIGN.CENTER
    p_link.space_before = Pt(6)

    # Save presentation
    output_filename = "Employee_Leave_Management_System_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as {output_filename}")

if __name__ == '__main__':
    create_deck()
